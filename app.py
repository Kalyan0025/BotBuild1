# --- Imports ---------------------------------------------------------------
import io
import re
import json
import streamlit as st
from google import genai
from google.genai import types
import PyPDF2
import docx

# --- Secrets ---------------------------------------------------------------
API_KEY = st.secrets.get("GEMINI_API_KEY", "")
if not API_KEY:
    st.warning("⚠️ 'GEMINI_API_KEY' is not set in st.secrets. Add it before deploying.")
    st.stop()

client = genai.Client(api_key=API_KEY)

# --- Identity / System Instructions ----------------------------------------
def load_identity() -> str:
    try:
        with open("identity.txt") as f:
            return f.read()
    except FileNotFoundError:
        return (
            "You are ReadysetRole. Output Overleaf-ready LaTeX for resume and cover letter. "
            "No fabrication. Use <METRIC_TBD> for unknown metrics. Resume ends [END_LATEX_RESUME], "
            "cover letter ends [END_LATEX_COVER]."
        )

SYSTEM_INSTRUCTIONS = load_identity()

# --- Helpers ---------------------------------------------------------------
def parse_resume_file(uploaded_file) -> str:
    """Extract text from PDF, DOCX, or TXT file."""
    try:
        if uploaded_file.type == "application/pdf":
            reader = PyPDF2.PdfReader(io.BytesIO(uploaded_file.read()))
            text = ""
            for p in reader.pages:
                try:
                    text += p.extract_text() or ""
                except Exception:
                    pass
            return text
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            docf = docx.Document(io.BytesIO(uploaded_file.read()))
            return "\n".join(para.text for para in docf.paragraphs)
        else:
            return uploaded_file.read().decode("utf-8", errors="ignore")
    except Exception as e:
        st.error(f"Error parsing file: {e}")
        return ""

def call_gemini(prompt: str, temperature: float = 0.5) -> str:
    """Call Gemini with system instructions."""
    try:
        cfg = types.GenerateContentConfig(
            system_instruction=SYSTEM_INSTRUCTIONS,
            temperature=temperature,
            max_output_tokens=8000,
        )
        resp = client.models.generate_content(
            model="gemini-2.0-flash-exp",
            contents=[types.Content(parts=[types.Part(text=prompt)])],
            config=cfg,
        )
        return resp.text or ""
    except Exception as e:
        st.error(f"Gemini API error: {e}")
        return ""

def extract_json(text: str):
    """Extract first valid JSON object/array from text."""
    try:
        m = re.search(r'```(?:json)?\s*(\{.*?\}|\[.*?\])\s*```', text, re.DOTALL)
        if m:
            return json.loads(m.group(1))
        m = re.search(r'(\{.*?\}|\[.*?\])', text, re.DOTALL)
        if m:
            return json.loads(m.group(1))
        return json.loads(text)
    except Exception:
        return None

# --- Prompts ---------------------------------------------------------------
SCORE_PROMPT_TMPL = """Return ONLY a JSON object with the fields below (0–100 integers).

Analyze this resume against the job description.

RESUME:
{resume}

JOB DESCRIPTION:
{jd}

Return:
{{
  "overall_score": 0,
  "skills_fit": 0,
  "experience_fit": 0,
  "education_fit": 0,
  "ats_keywords_coverage": 0
}}
"""

TAILOR_LATEX_PROMPT_TMPL = r"""Using the resume and JD below, output an Overleaf-ready ATS-safe LaTeX resume
that follows EXACTLY this structure (fill fields; omit sections if empty). Do not add tables/icons/graphics.
Use bullet style: Action → What → How/Tools → Impact. If impact metric is unknown, use <METRIC_TBD>.
The output MUST be ONLY the LaTeX code and MUST end with [END_LATEX_RESUME].

=== TEMPLATE TO FOLLOW ===
\documentclass[10pt]{article}
\usepackage[margin=0.5in]{geometry}
\usepackage[hidelinks]{hyperref}
\usepackage{enumitem}
\usepackage{titlesec}
\setlength{\parindent}{0pt}
\setlength{\parskip}{0pt}
\pagestyle{empty}

\titlespacing*{\section}{0pt}{4pt}{2pt}
\setlist[itemize]{noitemsep, topsep=2pt, leftmargin=1.2em}

\begin{document}

{\Large \bfseries %(name)s} \\[6pt]
%(location)s — %(phone)s — \href{mailto:%(email)s}{%(email)s} — %
\href{%(portfolio_url)s}{%(portfolio_label)s} — %
\href{%(linkedin_url)s}{%(linkedin_label)s}

\vspace{4pt}\hrule\vspace{4pt}

\section*{Summary}
%(summary)s

\vspace{4pt}\hrule\vspace{4pt}

\section*{Education}
%(education)s

\vspace{4pt}\hrule\vspace{4pt}

\section*{Professional Experience}
%(experience)s

\vspace{4pt}\hrule\vspace{4pt}

\section*{Selected Projects}
%(projects)s

\vspace{4pt}\hrule\vspace{4pt}

\section*{Skills}
%(skills)s

%(certs_block)s

\end{document}
=== END TEMPLATE ===

# Header Fields (fill exactly into the placeholders):
NAME: {name}
LOCATION: {location}
PHONE: {phone}
EMAIL: {email}
PORTFOLIO_URL: {portfolio_url}
PORTFOLIO_LABEL: {portfolio_label}
LINKEDIN_URL: {linkedin_url}
LINKEDIN_LABEL: {linkedin_label}

# Build content for Summary, Education, Experience, Projects, Skills, Certifications from the resume.
# Integrate JD keywords only where there is evidence in the resume. No fabrication. Use <METRIC_TBD> if needed.

RESUME:
{resume}

JOB DESCRIPTION:
{jd}

OUTPUT ONLY THE LATEX CODE. END WITH [END_LATEX_RESUME].
"""

COVER_LETTER_LATEX_PROMPT_TMPL = r"""Write a concise LaTeX cover letter (180–250 words) that CONTINUES the same resume–JD context.
It must use this exact LaTeX format and end with [END_LATEX_COVER].
Ground claims ONLY in the tailored resume; you may incorporate these user notes if provided.
No fabrication.

TEMPLATE:
\input{setup/preamble.tex}
\input{setup/macros.tex}

\begin{document}
\name{{%(name)s}}{{%(sender_title)s}}
\receiver{{%(receiver)s}}

\para{{Dear %(greeting)s,}}

\para{{[PARAGRAPH 1 — Connect resume strengths to the company/role using JD keywords that match the resume evidence.]}}

\para{{[PARAGRAPH 2 — Demonstrate 2–3 relevant achievements from the tailored resume (use tools/techniques).]}}

\para{{[PARAGRAPH 3 — Motivation, cultural fit, and a confident close with availability.] }}

\bottom{{%(sender_city)s}}{{%(sender_phone)s}}{{%(sender_email)s}}

\end{document}

CONTEXT (TAILORED RESUME — factual source of truth):
{tailored_resume}

COMPANY: {company}
ROLE: {role}
USER NOTES (optional): {notes}

OUTPUT ONLY THE LATEX CODE. END WITH [END_LATEX_COVER].
"""

# --- Session State ---------------------------------------------------------
if 'master_resume' not in st.session_state:
    st.session_state.master_resume = None
if 'master_resume_name' not in st.session_state:
    st.session_state.master_resume_name = None
if 'current_jd' not in st.session_state:
    st.session_state.current_jd = None
if 'scores' not in st.session_state:
    st.session_state.scores = {}
if 'tailored_latex' not in st.session_state:
    st.session_state.tailored_latex = None

# --- Page ------------------------------------------------------------------
st.set_page_config(page_title="ReadysetRole — LaTeX ATS Tailor", page_icon="⚡", layout="wide")
st.markdown("<h1 style='text-align:center'>⚡ ReadysetRole — LaTeX ATS Tailor</h1>", unsafe_allow_html=True)
st.caption("Upload resume + JD → QuickScore → Generate LaTeX Resume → (optional) LaTeX Cover Letter")

# --- Uploads ---------------------------------------------------------------
st.subheader("📤 Upload")
c1, c2 = st.columns(2)

with c1:
    up_res = st.file_uploader("Master Resume", type=["pdf", "docx", "txt"], key="resume_uploader")
    if up_res and st.session_state.master_resume is None:
        txt = parse_resume_file(up_res)
        if txt.strip():
            st.session_state.master_resume = txt
            st.session_state.master_resume_name = up_res.name
            st.success("✅ Resume loaded")
    if st.session_state.master_resume_name:
        st.info(f"📄 {st.session_state.master_resume_name}")

with c2:
    jd_mode = st.radio("Job Description", ["Paste Text", "Upload File"], horizontal=True)
    if jd_mode == "Paste Text":
        jd_txt = st.text_area("Paste JD", height=160, label_visibility="collapsed", key="jd_textarea")
        if st.button("Compute % Match (QuickScore)", use_container_width=True):
            if jd_txt and st.session_state.master_resume:
                st.session_state.current_jd = jd_txt
            else:
                st.warning("Please upload both resume and JD first.")
    else:
        up_jd = st.file_uploader("Upload JD", type=["pdf", "docx", "txt"], key="jd_uploader")
        if up_jd and st.button("Compute % Match (QuickScore)", use_container_width=True):
            jd_txt = parse_resume_file(up_jd)
            if jd_txt and st.session_state.master_resume:
                st.session_state.current_jd = jd_txt
            else:
                st.warning("Please upload both resume and JD first.")

st.divider()

# --- QuickScore ------------------------------------------------------------
if st.session_state.master_resume and st.session_state.current_jd:
    with st.spinner("Scoring..."):
        prompt = SCORE_PROMPT_TMPL.format(
            resume=st.session_state.master_resume[:6000],
            jd=st.session_state.current_jd[:6000]
        )
        raw = call_gemini(prompt, temperature=0.2)
        scores = extract_json(raw) or {}
        st.session_state.scores = scores

    s = st.session_state.scores or {}
    A, B, C, D, E = st.columns(5)
    A.metric("Overall Match", f"{int(s.get('overall_score', 0))}%")
    B.metric("Skills Fit", f"{int(s.get('skills_fit', 0))}%")
    C.metric("Experience Fit", f"{int(s.get('experience_fit', 0))}%")
    D.metric("Education Fit", f"{int(s.get('education_fit', 0))}%")
    E.metric("ATS Keywords", f"{int(s.get('ats_keywords_coverage', 0))}%")

    st.success("Next: fill header fields and click **Generate LaTeX Resume**")

    # --- Header fields for LaTeX resume -----------------------------------
    with st.form("latex_header_form"):
        st.subheader("👤 Header (Resume LaTeX)")
        name = st.text_input("Name", value="Your Name")
        location = st.text_input("Location", value="City, ST")
        phone = st.text_input("Phone", value="(000) 000-0000")
        email = st.text_input("Email", value="you@example.com")
        portfolio_url = st.text_input("Portfolio URL", value="https://example.com")
        portfolio_label = st.text_input("Portfolio Label", value="example.com")
        linkedin_url = st.text_input("LinkedIn URL", value="https://www.linkedin.com/in/your-handle/")
        linkedin_label = st.text_input("LinkedIn Label", value="linkedin.com/in/your-handle")
        generate_resume = st.form_submit_button("🎯 Generate LaTeX Resume", type="primary", use_container_width=True)

    if generate_resume:
        with st.spinner("Tailoring LaTeX resume..."):
            tailor_prompt = TAILOR_LATEX_PROMPT_TMPL.format(
                name=name,
                location=location,
                phone=phone,
                email=email,
                portfolio_url=portfolio_url,
                portfolio_label=portfolio_label,
                linkedin_url=linkedin_url,
                linkedin_label=linkedin_label,
                resume=st.session_state.master_resume,
                jd=st.session_state.current_jd
            )
            latex_resume = call_gemini(tailor_prompt, temperature=0.6)
            st.session_state.tailored_latex = latex_resume

# --- Tailored Resume (LaTeX) -----------------------------------------------
if st.session_state.tailored_latex:
    st.subheader("📄 Overleaf-Ready LaTeX Resume")
    st.code(st.session_state.tailored_latex, language="latex")

    st.divider()
    st.subheader("✉️ Optional: LaTeX Cover Letter")

    with st.form("cover_form"):
        company = st.text_input("Company", value="")
        role = st.text_input("Role / Position", value="")
        receiver = st.text_input("Receiver (e.g., Hiring Manager \\\\ Company)", value="Hiring Manager")
        greeting = st.text_input("Greeting (Dear ___,)", value="Hiring Manager")
        sender_title = st.text_input("Sender Title (under \\name{})", value="Applicant")
        sender_city = st.text_input("Sender City", value="City, ST")
        sender_phone = st.text_input("Sender Phone", value="(000) 000-0000")
        sender_email = st.text_input("Sender Email", value="you@example.com")
        notes = st.text_area("Optional notes to emphasize (kept factual)", value="", height=100)
        gen_cover = st.form_submit_button("Generate LaTeX Cover Letter", use_container_width=True)

    if gen_cover:
        with st.spinner("Drafting LaTeX cover letter..."):
            cl_prompt = COVER_LETTER_LATEX_PROMPT_TMPL.format(
                tailored_resume=st.session_state.tailored_latex,
                company=company,
                role=role,
                notes=notes
            ) % {
                "name": name,
                "sender_title": sender_title,
                "receiver": receiver,
                "greeting": greeting,
                "sender_city": sender_city,
                "sender_phone": sender_phone,
                "sender_email": sender_email,
            }
            latex_cover = call_gemini(cl_prompt, temperature=0.6)
            st.code(latex_cover, language="latex")

st.caption("ReadysetRole — LaTeX-first ATS Tailoring (no fabrication)")
